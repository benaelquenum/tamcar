"""Moteur de simulation du classeur v2 : reproduit les formules du classeur (contrôlé contre Excel) pour balayer vite
taux d'atteinte, nombre de véhicules de tiers, seuils. Les séries de base viennent d'un classeur recalculé par Excel
(taux 100 %, aucun véhicule de tiers)."""
import math
import pathlib
import subprocess
import tempfile

import openpyxl

import modele_v2 as MV

ICI = pathlib.Path(__file__).resolve().parent
PS1 = ICI / "recalc.ps1"
LEVEE = 15_000_000
JOURS = 26
COLS = range(7, 44)   # mois 0 à 36


def recalculer(chemin):
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(PS1), "-Path", str(chemin)],
                       capture_output=True, text=True, timeout=1800)
    assert "OK" in r.stdout, (r.stdout, r.stderr)


def lire_mois(ws, r):
    return [ws.cell(r, c).value or 0 for c in COLS]


def arrondi(x):
    return math.floor(x + 0.5)


class Base:
    """Séries mensuelles (indice = mois 0..36) du classeur à 100 %, sans véhicule de tiers."""

    def __init__(self, dossier=None):
        dossier = pathlib.Path(dossier or tempfile.mkdtemp(prefix="tamcar_base_"))
        dossier.mkdir(parents=True, exist_ok=True)
        f = dossier / "base0.xlsx"
        MV.construire(f, surcharges={"B94": 1, "B95": 0, "B96": 0, "B97": 0, "B98": 0})
        recalculer(f)
        wb = openpyxl.load_workbook(f, data_only=True)
        M, H = wb["Mensuel"], wb["Hypothèses"]
        g = lambda r: lire_mois(M, r)
        # recettes BRUTES de la cotisation TamAssur (ligne 62 + ligne 102) : la cotisation est fixe, elle ne suit pas le taux d'atteinte
        self.occ = g(102)
        self.rec = [a + b for a, b in zip(g(62), self.occ)]
        self.ch = g(54)
        self.achat = g(45)
        self.ops = g(33)
        # ligne 38 = postes vacants + manque au rachat NSIA (ligne 105)
        self.nsia = g(105)
        self.vac = [a - b for a, b in zip(g(38), self.nsia)]
        self.entree = g(37)
        self.anatt, self.imf, self.patente = g(34), g(35), g(36)
        self.amort = [M.cell(73, c).value for c in (3, 4, 5)]
        self.frais_demarrage = M.cell(74, 2).value + M.cell(74, 3).value
        self.gadgets_gestes = [M.cell(77, c).value for c in (3, 4, 5)]
        self.demarrage = M["B20"].value
        self.h = {k: H[k].value for k in ("B45", "B46", "B57", "B58", "B59", "B60", "B61", "B63", "B43", "B44", "B47", "B80", "B81", "B82", "B101", "B54", "B103", "B104", "B105")}
        self.is_taux, self.is_exo = H["B88"].value, H["B89"].value
        wb.close()

    def to_dict(self):
        return dict(self.__dict__)

    @classmethod
    def from_dict(cls, d):
        o = cls.__new__(cls)
        o.__dict__.update(d)
        return o


def simuler(base, p=1.0, ess=0, conf=0, tri=0, moto=0, debut=4, fin=12, part_tm=0.25, delta_part=0.0):
    """Séries mensuelles (mois 0..36) : recettes, charges, ebe, flux, caisse. `delta_part` : points de part de TamCar en plus sur les voitures (Afrik Group et tiers)."""
    h = base.h
    jours, plan = h["B47"], dict(ess=h["B60"], conf=h["B61"], tri=h["B45"], moto=h["B46"])
    n_tiers = dict(ess=ess, conf=conf, tri=tri, moto=moto)
    cum = {k: [0] * 37 for k in n_tiers}
    for k, n in n_tiers.items():
        for m in range(1, 37):
            cum[k][m] = arrondi(n * max(0.0, min(1.0, (m - debut + 1) / (fin - debut + 1))))
    vacance_unit = h["B81"] / 12 * h["B82"] * h["B80"]
    cap_ops = h["B44"]
    out = dict(rec=[0.0] * 37, ch=[0.0] * 37, ebe=[0.0] * 37, flux=[0.0] * 37, caisse=[0.0] * 37, nsia=[0.0] * 37)
    caisse = LEVEE
    # mois 0
    out["ch"][0] = base.ch[0]
    caisse -= base.ch[0]
    out["flux"][0] = -base.ch[0]
    out["caisse"][0] = caisse
    for m in range(1, 37):
        rec = p * base.rec[m] - base.occ[m]
        if delta_part and m >= h["B59"]:
            rec += p * jours * delta_part * (h["B57"] * plan["ess"] + h["B58"] * plan["conf"]) * (0.5 if m == h["B59"] else 1.0)
        vol_tiers_jour = 0.0
        nsia_tiers = 0.0
        for k in n_tiers:
            c, c0 = cum[k][m], cum[k][m - 1]
            eff = c - 0.5 * (c - c0)
            part = (h["B63"] + delta_part) if k in ("ess", "conf") else part_tm
            rec += p * jours * plan[k] * eff * part - h["B103"] * jours * eff      # part de TamCar, moins la cotisation TamAssur (fixe)
            vol_tiers_jour += cum[k][m] * plan[k]
            # manque au rachat NSIA à l'ouverture du premier palier : 12 mois pour une moto, 24 mois pour les autres
            fenetre, manque = (12, h["B104"]) if k == "moto" else (24, h["B105"])
            if m - fenetre >= 1:
                nsia_tiers += manque * (cum[k][m - fenetre] - cum[k][m - fenetre - 1])
        d0 = base.ops[m] / (h["B43"] * jours)                 # volume journalier de base (plafond non atteint)
        ops = min(h["B43"] * p * (d0 + vol_tiers_jour) * jours, cap_ops)
        nb0 = base.vac[m] / vacance_unit
        vac = (nb0 + sum(cum[k][m] for k in cum)) * vacance_unit
        arrivees = sum(cum[k][m] - cum[k][m - 1] for k in cum)
        entree = h["B101"] * (base.entree[m] / h["B101"] + arrivees)
        ch = base.ch[m] - base.ops[m] + ops - base.vac[m] + vac - base.entree[m] + entree + nsia_tiers
        out["rec"][m] = rec
        out["nsia"][m] = base.nsia[m] + nsia_tiers
        out["ch"][m] = ch
        out["ebe"][m] = rec - (ch - base.achat[m])
        out["flux"][m] = rec - ch
        caisse += rec - ch
        out["caisse"][m] = caisse
    return out


def annuel(serie, cle):
    return [sum(serie[cle][12 * i + 1:12 * i + 13]) for i in range(3)]


def resultats(base, s):
    ebe = annuel(s, "ebe")
    res_exploit = [ebe[0] - base.amort[0] - base.frais_demarrage, ebe[1] - base.amort[1], ebe[2] - base.amort[2]]
    impot = []
    cum_perte = 0.0
    for i, r in enumerate(res_exploit):
        if i + 1 <= base.is_exo:
            impot.append(0.0)
        else:
            impot.append(base.is_taux * max(0.0, r + min(0.0, sum(res_exploit[:i]))))
    net = [r - t for r, t in zip(res_exploit, impot)]
    net_apres = [n - g for n, g in zip(net, base.gadgets_gestes)]
    caisse = s["caisse"]
    return dict(ebe=ebe, net=net, net_apres=net_apres, caisse_fin=[caisse[12], caisse[24], caisse[36]],
                point_bas=min(caisse), mois_negatifs=sum(1 for v in caisse if v < 0),
                premier_mois_negatif=next((m for m, v in enumerate(caisse) if v < 0), None))


def chaque_mois(base, s, a_partir=13):
    """Excédent d'exploitation ET flux de caisse >= 0 chaque mois à partir du mois indiqué, charges annuelles (ANATT, impôt minimum,
    patente) lissées en réserve mensuelle."""
    # le manque au rachat NSIA (ponctuel, à l'ouverture du premier palier de retrait) est mis en réserve comme les charges annuelles
    lissage = (sum(base.anatt) + sum(base.imf) + sum(base.patente) + sum(s["nsia"][13:])) / 24   # réserve mensuelle sur les 24 mois où ces charges tombent (mois 13 à 36)
    for m in range(a_partir, 37):
        reserve = (base.anatt[m] + base.imf[m] + base.patente[m] + s["nsia"][m]) - lissage   # on remet la charge réelle du mois et on retire la réserve mensuelle
        if s["ebe"][m] + reserve < -0.5 or s["flux"][m] + reserve < -0.5:
            return False
    return True


def minimal(base, p, genre, a_partir=13, limite=60, **kw):
    for n in range(limite):
        s = simuler(base, p, **{genre: n}, **kw)
        if chaque_mois(base, s, a_partir):
            return n
    return None


def seuil(base, f, bas=0.2, haut=1.6, **kw):
    """Plus petit taux d'atteinte pour lequel f(resultats) >= 0."""
    for _ in range(60):
        mid = (bas + haut) / 2
        if f(resultats(base, simuler(base, mid, **kw))) >= 0:
            haut = mid
        else:
            bas = mid
    return haut
