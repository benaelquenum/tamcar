"""Classeur TamCar « Résultat complet » — version 2 (décisions du 2026-10-11, du 2026-10-18 et du 2026-10-20).

Part du classeur v1 archivé (archives/2026-10-11_classeur_resultat_complet_v1) et applique :
  - fonds de rachat ramené de 10 % à 5 % : les 5 points retirés vont à TamCar -> 25 % du volume sur les voitures (20 % + 5 %), 55 % sur un tricycle de TamCar ;
  - durées de contrat : Essentiel 36 mois, Confort 48 mois (VIP 60 mois, non compté) ;
  - plus d'avance de démarrage remboursable : 30 000 F de frais d'entrée par véhicule déployé (charge) ;
  - GPS / tracker, assurance et entretien des voitures à la charge du propriétaire ;
  - un taux d'atteinte des planchers (levier de pilotage) et le recrutement de véhicules de tiers (3 Essentiel + 3 Confort au minimum) ;
  - (2026-10-20) planchers de recettes corrigés : objectif du jour « tout compris » (versement + TamAssur 1 000 F) ÷ 60 %, soit moto 5 850 F, tricycle 10 000 F,
    Essentiel 13 700 F, Confort 16 700 F ; la cotisation TamAssur (1 000 F par jour et par véhicule sous contrat) est financée par la part de TamCar, plus par
    le conducteur en plus de son versement ; manque au rachat NSIA compté à l'ouverture du premier palier de retrait (12 mois pour la moto, 24 mois pour les autres).

Les formules sont celles de v1, complétées : rien n'est figé dans les feuilles de calcul (sauf la feuille « Seuils », qui garde le résultat
des simulations avec sa date).
"""
import math
import pathlib
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter as L

ICI = pathlib.Path(__file__).resolve().parent
RACINE = ICI.parent.parent
V1 = RACINE / "archives" / "2026-10-11_classeur_resultat_complet_v1" / "TamCar_resultat_complet.xlsx"
V2 = RACINE / "TamCar_resultat_complet_v2.xlsx"

HY = "Hypothèses!"
P, N_ESS, N_CONF, N_TRI, N_MOTO = (f"{HY}$B${r}" for r in (94, 95, 96, 97, 98))
DEBUT, FIN, ENTREE, PART_TM = (f"{HY}$B${r}" for r in (99, 100, 101, 102))

# valeurs par défaut du classeur v2 (cellule -> valeur) ; les scénarios les surchargent
DEFAUT = {"B94": 1, "B95": 3, "B96": 3, "B97": 0, "B98": 0, "B99": 4, "B100": 12, "B101": 30000, "B102": 0.25}


def _remplacer(valeur, ancien, nouveau, tag):
    assert ancien in valeur, (tag, ancien[:70], valeur[:160])
    return valeur.replace(ancien, nouveau, 1)


def construire(dest=V2, surcharges=None, source=V1):
    """Écrit le classeur v2 (formules, sans valeurs en cache) ; `surcharges` : {'B94': 0.7, ...} sur la feuille Hypothèses."""
    wb = openpyxl.load_workbook(source)
    H, M = wb["Hypothèses"], wb["Mensuel"]

    # ------------------------------------------------------------------ Hypothèses : décisions
    def note(r, texte):
        H.cell(r, 3).value = texte

    H["B62"], H["B63"] = 0.25, 0.25
    H["A62"] = "Part de TamCar sur le volume des voitures, première année"
    note(62, "20 % plateforme + 5 % : le fonds de rachat passe de 10 % à 5 % et les 5 points retirés vont à TamCar (décision du 2026-10-18) ; le propriétaire garde 30 % et reçoit le fonds de 5 % à la cession")
    H["A63"] = "Part de TamCar sur le volume des voitures, jusqu'à la cession"
    note(63, "Idem : 20 % + 5 %")
    H["B65"], H["B66"] = 36, 48
    note(65, "Terence (2026-10-18) : 3 ans pour l'Essentiel (Confort 4 ans, VIP 5 ans, non compté)")
    note(66, "Terence (2026-10-18) : 4 ans pour le Confort")
    H["A67"] = "Avance de démarrage du propriétaire, par voiture (supprimée)"
    H["B67"] = 0
    note(67, "Supprimée (décision du 2026-10-11) : TamCar avance les frais d'entrée (ligne 101), aucun remboursement au mois 13")
    H["A68"] = "Assurance, trackers GPS, entretien des voitures"
    note(68, "À la charge du propriétaire, y compris le GPS / tracker (confirmé par Terence le 2026-10-18)")
    H["B71"], H["B72"] = 0.55, 0.55
    H["A71"] = "Part de TamCar sur le volume d'un tricycle, première année"
    note(71, "30 % propriétaire (TamCar) + 20 % plateforme + 5 % fonds de rachat qui revient à TamCar ; l'autre moitié du fonds (5 %) est le fonds du conducteur")
    note(72, "Idem : 55 %")

    # ------------------------------------------------------------------ Hypothèses : spots vidéo par IA dans le marketing de lancement (2026-10-23)
    note(27, "Comprend la création de spots vidéo par IA (Higgsfield) : pilote 39 200 F (forfait Plus, 49 $, 1 mois), puis 309 600 F au plus (forfait Ultra, 3 × 129 $) si le pilote convainc, dollar à 800 F ; le solde couvre la photo et le sponsoring ; la diffusion payante n'est pas budgétée")

    # ------------------------------------------------------------------ Hypothèses : tricycle cédé à 24 mois (décision du 2026-10-21)
    H["B74"] = 24
    note(74, "Terence (2026-10-21) : 2 ans pour le tricycle (garantie TamAssur de 500 000 F à 24 mois, qui est aussi le terme) ; assurance et tracker payés à l'achat puis au 1er anniversaire seulement ; 20 % de commission ensuite")

    # ------------------------------------------------------------------ Hypothèses : planchers corrigés (2026-10-20)
    H["B45"], H["B46"], H["B60"], H["B61"] = 10000, 5850, 13700, 16700
    note(45, "Objectif du jour du tricycle 6 000 F (versement 5 000 F + TamAssur 1 000 F) = 60 % du volume : 10 000 F (décision du 2026-10-20) ; sert au responsable des opérations")
    note(46, "Objectif du jour de la moto 3 500 F (versement 2 500 F + TamAssur 1 000 F) = 60 % du volume : 5 850 F (décision du 2026-10-20)")
    note(60, "Objectif du jour 8 200 F (versement 7 200 F + TamAssur 1 000 F) = 60 % du volume : 13 700 F ; le conducteur a déjà pris ses 40 % (décision du 2026-10-20)")
    note(61, "Objectif du jour 10 000 F (versement 9 000 F + TamAssur 1 000 F) = 60 % du volume : 16 700 F (décision du 2026-10-20)")
    H["A75"] = "Objectif du jour d'une moto de TamCar, tout compris (versement 2 500 F + TamAssur 1 000 F)"
    H["B75"] = 3500
    note(75, "Terence (2026-10-15) : 3 500 F par jour, 26 jours par mois, pendant 12 mois ; la cotisation TamAssur de 1 000 F est déduite plus bas (ligne 103) ; plus si l'app rapporte plus")
    H["A79"] = "Épargne NSIA (TamAssur) : cotisation, postes vacants, manque au rachat"
    note(80, "Cotisation payée à NSIA chaque jour par TamCar pour chaque conducteur sous contrat (ligne 103) ; le conducteur ne la verse PAS en plus : elle est comprise dans son objectif du jour, donc dans les 60 % du volume partagés entre propriétaire, fonds de rachat et TamCar")
    note(82, "Hypothèse de travail à confirmer ; TamCar paie alors la cotisation de sa poche (poste vacant : aucune course, donc aucune part)")

    # ------------------------------------------------------------------ Hypothèses : leviers (lignes 93 à 102)
    H["A93"] = "Pilotage et recrutement (ajout du 2026-10-18)"
    for col in "ABC":
        H[f"{col}93"]._style = copy(H[f"{col}86"]._style)
    lignes = [
        (94, "Taux d'atteinte des planchers (100 % = planchers atteints)", 1, "0%", "Tout le volume (versements des motos compris) est multiplié par ce taux. Mettre 70 % pour le scénario de référence de sécurité ; voir la feuille Seuils"),
        (95, "Voitures Essentiel de tiers à recruter (en plus des 6 d'Afrik Group)", 3, "0", "Décision du 2026-10-18 : au minimum 3 Essentiel + 3 Confort de tiers ; mettre 0 pour les exclure"),
        (96, "Voitures Confort de tiers à recruter (en plus des 3 d'Afrik Group)", 3, "0", "Idem"),
        (97, "Tricycles de tiers à recruter", 0, "0", "Option : un tricycle de tiers rapporte 25 % du volume plancher à TamCar (hypothèse à confirmer)"),
        (98, "Motos de tiers à recruter", 0, "0", "Option : une moto de tiers rapporte 25 % du volume plancher (hypothèse à confirmer) ; peu rentable"),
        (99, "Arrivée des véhicules de tiers : à partir du mois", 4, "0", "Les arrivées sont étalées régulièrement entre ce mois et le suivant"),
        (100, "… jusqu'au mois", 12, "0", "Le dernier véhicule arrive à ce mois"),
        (101, "Frais d'entrée par véhicule déployé (inspection, habillage, kit chauffeur)", 30000, '#,##0;\\(#,##0\\);\\-', "Décision du 2026-10-11 : TamCar dépense 30 000 F par véhicule ; compté à l'arrivée de chaque véhicule (15 au mois 1, puis les véhicules de tiers)"),
        (102, "Part de TamCar sur un tricycle ou une moto de tiers", 0.25, "0%", "Même partage que les voitures : 30 % propriétaire, 20 % plateforme, 5 % fonds du conducteur, 5 % fonds qui revient à TamCar"),
        (103, "Cotisation TamAssur par jour et par véhicule sous contrat, financée par la part de TamCar (F)", "=B80", '#,##0;\\(#,##0\\);\\-', "Reprend la ligne 80. Décision du 2026-10-20 : planchers « tout compris » ; la cotisation est prise sur la part de TamCar (25 % du volume pour une voiture), pas ajoutée au versement du conducteur. Déduite des recettes, ne varie pas avec le taux d'atteinte"),
        (104, "Manque au rachat NSIA à 12 mois, par moto (F)", 29000, '#,##0;\\(#,##0\\);\\-', "Estimation du 2026-10-15 : garantie 250 000 F contre un rachat net NSIA d'environ 221 000 F (85 % de l'épargne, pénalité 5 %). À confirmer par une simulation écrite de NSIA ; mettre 0 pour exclure"),
        (105, "Manque au rachat NSIA à 24 mois, par tricycle ou voiture (F)", 48000, '#,##0;\\(#,##0\\);\\-', "Estimation du 2026-10-15 : garantie 500 000 F contre un rachat net d'environ 452 000 F ; les paliers suivants (36, 48, 60 mois) sont supposés sans manque. À confirmer par NSIA ; mettre 0 pour exclure"),
    ]
    for r, label, val, fmt, nt in lignes:
        H.cell(r, 1, label)._style = copy(H["A31"]._style)
        c = H.cell(r, 2, val)
        c._style = copy(H["B31"]._style)
        c.number_format = fmt
        H.cell(r, 3, nt)._style = copy(H["C31"]._style)
    for cell, val in (surcharges or {}).items():
        H[cell] = val

    # ------------------------------------------------------------------ Mensuel
    cols = range(7, 44)  # G (mois 0) à AQ (mois 36)

    def ligne(r, label, fn, style_row=38, gras=False):
        M.cell(r, 1, label)._style = copy(M.cell(style_row, 1)._style)
        for c in cols:
            cell = M.cell(r, c, fn(c, L(c), L(c - 1)))
            cell._style = copy(M.cell(style_row, c)._style)
            cell.number_format = '#,##0;\\(#,##0\\);\\-'
        M.cell(r, 2, f"=G{r}")
        M.cell(r, 3, f"=SUM(H{r}:S{r})")
        M.cell(r, 4, f"=SUM(T{r}:AE{r})")
        M.cell(r, 5, f"=SUM(AF{r}:AQ{r})")
        M.cell(r, 6, f"=SUM(B{r}:E{r})")
        for c in range(2, 7):
            M.cell(r, c)._style = copy(M.cell(style_row, c)._style)
            M.cell(r, c).number_format = '#,##0;\\(#,##0\\);\\-'

    # section « véhicules de tiers » (style de la ligne 79)
    for c in range(1, 44):
        M.cell(86, c)._style = copy(M.cell(79, c)._style)
    M.cell(86, 1, "VÉHICULES DE TIERS À RECRUTER (leviers de volume ; déjà compris dans les recettes et les volumes ci-dessus)")

    def cumul(nom):
        return lambda c, k, j: "=0" if c == 7 else (
            f"=ROUND({nom}*MAX(0,MIN(1,({k}$4-{DEBUT}+1)/({FIN}-{DEBUT}+1))),0)")

    ligne(87, "Voitures Essentiel de tiers en service (cumul)", cumul(N_ESS))
    ligne(88, "Voitures Confort de tiers en service (cumul)", cumul(N_CONF))
    ligne(89, "Tricycles de tiers en service (cumul)", cumul(N_TRI))
    ligne(90, "Motos de tiers en service (cumul)", cumul(N_MOTO))
    ligne(91, "Véhicules de tiers arrivés ce mois", lambda c, k, j: "=0" if c == 7 else f"=({k}87-{j}87)+({k}88-{j}88)+({k}89-{j}89)+({k}90-{j}90)")

    def vol_tiers(rang, plancher):
        return lambda c, k, j: "=0" if c == 7 else (
            f"={P}*{HY}$B$47*{plancher}*({k}{rang}-0.5*({k}{rang}-{j}{rang}))")

    ligne(92, "Volume des Essentiel de tiers (courses, au taux d'atteinte retenu)", vol_tiers(87, f"{HY}$B$60"))
    ligne(93, "Volume des Confort de tiers", vol_tiers(88, f"{HY}$B$61"))
    ligne(94, "Volume des tricycles de tiers", vol_tiers(89, f"{HY}$B$45"))
    ligne(95, "Volume des motos de tiers", vol_tiers(90, f"{HY}$B$46"))
    ligne(96, "Part de TamCar sur les véhicules de tiers (déjà dans les lignes 58 à 61)",
          lambda c, k, j: f"=({k}92+{k}93)*{HY}$B$63+({k}94+{k}95)*{PART_TM}")
    for r in (87, 88, 89, 90, 91):
        for c in range(7, 44):
            M.cell(r, c).number_format = '0'

    # section « TamAssur » (décision du 2026-10-20) : véhicules sous contrat, cotisation prise sur la part de TamCar, manque au rachat NSIA
    for c in range(1, 44):
        M.cell(97, c)._style = copy(M.cell(79, c)._style)
    M.cell(97, 1, "TAMASSUR : véhicules sous contrat (équivalents en service), cotisation prise sur la part de TamCar, manque au rachat NSIA")
    B = lambda n: f"{HY}$B${n}"

    def age(k):
        return f"({k}$4-$G$4:$AQ$4)"

    ligne(98, "Tricycles sous contrat, équivalents en service (TamCar + tiers)",
          lambda c, k, j: "=0" if c == 7 else
          f"=SUMPRODUCT($G$7:$AQ$7*({age(k)}>=0)*(1-0.5*({age(k)}=0))*({age(k)}<{B(74)}))+({k}89-0.5*({k}89-{j}89))")
    ligne(99, "Motos sous contrat, équivalents en service (TamCar + tiers)",
          lambda c, k, j: "=0" if c == 7 else
          f"=SUMPRODUCT($G$8:$AQ$8*({age(k)}>=0)*(1-0.5*({age(k)}=0))*({age(k)}<{B(76)}))+({k}90-0.5*({k}90-{j}90))")
    ligne(100, "Voitures Essentiel sous contrat, équivalents en service (Afrik Group + tiers)",
          lambda c, k, j: "=0" if c == 7 else
          f"=IF({k}$4<{B(59)},0,{B(57)}*IF({k}$4={B(59)},0.5,1)*(({k}$4-{B(59)})<{B(65)}))+({k}87-0.5*({k}87-{j}87))")
    ligne(101, "Voitures Confort sous contrat, équivalents en service (Afrik Group + tiers)",
          lambda c, k, j: "=0" if c == 7 else
          f"=IF({k}$4<{B(59)},0,{B(58)}*IF({k}$4={B(59)},0.5,1)*(({k}$4-{B(59)})<{B(66)}))+({k}88-0.5*({k}88-{j}88))")
    ligne(102, "Cotisation TamAssur prise sur la part de TamCar (déjà déduite des lignes 58 à 61)",
          lambda c, k, j: "=0" if c == 7 else f"={B(103)}*{B(47)}*({k}98+{k}99+{k}100+{k}101)")
    ligne(103, "Motos arrivées ce mois (premier palier de retrait à 12 mois)",
          lambda c, k, j: "=0" if c == 7 else f"=N({k}8)+({k}90-{j}90)")
    ligne(104, "Tricycles et voitures arrivés ce mois (premier palier de retrait à 24 mois)",
          lambda c, k, j: "=0" if c == 7 else
          f"=N({k}7)+({k}89-{j}89)+IF({k}$4={B(59)},{B(57)}+{B(58)},0)+({k}87-{j}87)+({k}88-{j}88)")
    ligne(105, "Manque au rachat NSIA à l'ouverture du premier palier (compté dans la ligne 38)",
          lambda c, k, j: "=0" if c == 7 else
          f"={B(104)}*SUMPRODUCT($G$103:$AQ$103*($G$4:$AQ$4={k}$4-12))+{B(105)}*SUMPRODUCT($G$104:$AQ$104*($G$4:$AQ$4={k}$4-24))")
    for r in (98, 99, 100, 101):
        for c in range(7, 44):
            M.cell(r, c).number_format = '0.0'

    for c in cols:
        k = L(c)
        # responsable des opérations : taux d'atteinte et véhicules de tiers
        v = M.cell(33, c).value
        if isinstance(v, str) and v.startswith("=MIN("):
            v = _remplacer(v, f"{HY}$B$43*(", f"{HY}$B$43*{P}*(", f"M33 {k}")
            fin_somme = f"IF({k}$4>={HY}$B$59,{HY}$B$57*{HY}$B$60+{HY}$B$58*{HY}$B$61,0))"
            v = _remplacer(v, fin_somme, fin_somme[:-1] + f"+{k}87*{HY}$B$60+{k}88*{HY}$B$61+{k}89*{HY}$B$45+{k}90*{HY}$B$46)", f"M33b {k}")
            M.cell(33, c).value = v
        # tricycle cédé à 24 mois : plus d'assurance ni de tracker au 2e anniversaire (mois 25)
        for r_ in (48, 50):
            v = M.cell(r_, c).value
            if isinstance(v, str) and "-24)*$G$7:$AQ$7)" in v:
                M.cell(r_, c).value = _remplacer(v, f"+SUMPRODUCT(($G$4:$AQ$4={k}$4-24)*$G$7:$AQ$7)", "", f"M{r_} {k}")
        # frais d'entrée (remplace l'avance remboursée)
        M.cell(37, c).value = (f"={ENTREE}*(N({k}7)+N({k}8)+IF({k}$4={HY}$B$59,{HY}$B$57+{HY}$B$58,0)+{k}91)"
                               if c > 7 else "=0")
        # postes vacants NSIA : + véhicules de tiers ; + manque au rachat NSIA (ligne 105)
        v = M.cell(38, c).value
        if isinstance(v, str) and v.startswith("=("):
            prem = f"IF({k}$4>={HY}$B$59,{HY}$B$57+{HY}$B$58,0))"
            v = _remplacer(v, prem, prem[:-1] + f"+{k}87+{k}88+{k}89+{k}90)", f"M38 {k}")
            M.cell(38, c).value = v + f"+{k}105"
        # recettes : taux d'atteinte + véhicules de tiers
        tam = lambda ligne_eq: f"-{HY}$B$103*{HY}$B$47*{k}{ligne_eq}"      # cotisation TamAssur prise sur la part de TamCar (fixe, hors taux d'atteinte)
        for r, extra in ((58, f"+{k}94*{PART_TM}{tam(98)}"), (59, f"+{k}95*{PART_TM}{tam(99)}"),
                         (60, f"+{k}92*{HY}$B$63{tam(100)}"), (61, f"+{k}93*{HY}$B$63{tam(101)}")):
            v = M.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                M.cell(r, c).value = f"=({v[1:]})*{P}{extra}"
        # volumes
        for r, extra in ((80, f"+{k}94"), (81, f"+{k}95"), (82, f"+{k}92"), (83, f"+{k}93")):
            v = M.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                M.cell(r, c).value = f"=({v[1:]})*{P}{extra}"

    # libellés
    M["A37"] = "Frais d'entrée des véhicules déployés (inspection, habillage, kit : 30 000 F par véhicule)"
    M["A38"] = "Épargne NSIA (TamAssur) : postes vacants et manque au rachat à l'ouverture du premier palier"
    M["A33"] = "Responsable des opérations (3 % du volume, plafonné)"
    M["A80"] = "Volume des tricycles (TamCar + tiers)"
    M["A81"] = "Volume des motos (TamCar + tiers)"
    M["A82"] = "Volume des voitures Essentiel (Afrik Group + tiers)"
    M["A83"] = "Volume des voitures Confort (Afrik Group + tiers)"
    M["A79"] = "VOLUME DE COURSES (au taux d'atteinte retenu ; sert à répartir le responsable des opérations ; mois de lancement compté pour moitié)"
    for r, t in ((58, "Tricycles : part de TamCar sur le volume, nette de la cotisation TamAssur (TamCar + tiers)"),
                 (59, "Motos : objectif du jour puis commission après cession, net de la cotisation TamAssur (TamCar + tiers)"),
                 (60, "Voitures Essentiel : part de TamCar sur le volume, nette de la cotisation TamAssur (Afrik Group + tiers)"),
                 (61, "Voitures Confort : part de TamCar sur le volume, nette de la cotisation TamAssur (Afrik Group + tiers)")):
        M.cell(r, 1).value = t

    # ------------------------------------------------------------------ autres feuilles : libellés et notes
    ren = {"Voitures Essentiel (6)": "Voitures Essentiel (Afrik Group + tiers)",
           "Voitures Confort (3)": "Voitures Confort (Afrik Group + tiers)",
           "Motos (5)": "Motos (5 de TamCar + tiers)",
           "Tricycle (1)": "Tricycle (1 de TamCar + tiers)"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    for a, b in ren.items():
                        if cell.value == a:
                            cell.value = b
    wb["Synthèse"]["A30"] = "Frais d'entrée des véhicules déployés (30 000 F par véhicule)"
    wb["Synthèse"]["A31"] = "Épargne NSIA (TamAssur) : postes vacants et manque au rachat"
    wb["Synthèse"]["A73"] = ("Épargne NSIA (TamAssur) : la cotisation de 1 000 F par jour est comprise dans l'objectif du jour du conducteur, donc dans les 60 % du volume qui reviennent "
                             "au propriétaire, au fonds de rachat et à TamCar : TamCar la finance sur sa part (déduite des recettes). S'y ajoutent le coût des postes vacants et le manque au "
                             "rachat NSIA à l'ouverture du premier palier de retrait (estimations des lignes 104 et 105 des Hypothèses, à confirmer par NSIA).")
    wb["Charges"]["A58"] = ("Non inclus : frais FeexPay (1,7 % payés par le client), frais Mobile Money des retraits chauffeur, forfait internet des chauffeurs, charges sociales, entretien (conducteur). "
                            "La cotisation TamAssur est déduite des recettes ; le manque au rachat NSIA est compté dans l'épargne NSIA (estimation).")
    wb["Recettes"]["A26"] = "Moto : objectif du jour (dont TamAssur 1 000 F)"

    Mg = wb["Marge"]
    for bloc in (15, 30, 45):
        deb, fin = {15: ("H", "S"), 30: ("T", "AE"), 45: ("AF", "AQ")}[bloc]
        Mg[f"B{bloc}"] = f"=AVERAGE(Mensuel!{deb}10:{fin}10)+AVERAGE(Mensuel!{deb}90:{fin}90)"
        Mg[f"C{bloc}"] = f"=AVERAGE(Mensuel!{deb}9:{fin}9)+AVERAGE(Mensuel!{deb}89:{fin}89)"
        Mg[f"D{bloc}"] = f"={HY}$B$57+AVERAGE(Mensuel!{deb}87:{fin}87)"
        Mg[f"E{bloc}"] = f"={HY}$B$58+AVERAGE(Mensuel!{deb}88:{fin}88)"

    R = wb["Résultat"]
    R["A37"] = ("Lecture : plus d'avance de démarrage remboursable (décision du 2026-10-11) ; les frais d'entrée (30 000 F par véhicule déployé, "
                "15 véhicules au mois 1 puis les véhicules de tiers) sont une charge de fonctionnement de l'année d'arrivée.")
    D = wb["Démarrage"]
    D["A28"] = ("Les voitures (6 Essentiel, 3 Confort) appartiennent à Afrik Group : assurance, trackers GPS et entretien à la charge du propriétaire. "
                "TamCar avance 30 000 F de frais d'entrée par véhicule (pas d'avance remboursable).")

    S = wb["Résumé"]
    S["A2"] = ("Flotte : 1 tricycle + 5 motos de TamCar + 9 voitures d'Afrik Group (6 Essentiel, 3 Confort) + au minimum 6 voitures de tiers (3 Essentiel, 3 Confort) "
               "recrutées entre le mois 4 et le mois 12. Recettes aux planchers de Mayeul, au taux d'atteinte de la feuille Hypothèses (ligne 94). Levée de 15 M. Lancement en janvier 2027 (mois 1).")
    S["A27"] = ("Recettes aux planchers : seules les courses jusqu'au plancher de Mayeul sont comptées. Tout volume au-dessus du plancher s'ajoute aux recettes. "
                "Le premier mois de lancement (et le mois d'arrivée d'un véhicule) compte pour moitié. Part de TamCar : 25 % sur les voitures, 55 % sur le tricycle de TamCar, "
                "objectif du jour de 3 500 F sur les motos de TamCar (décision du 2026-10-18). Planchers corrigés le 2026-10-20 : volume = objectif ÷ 60 % ; la cotisation TamAssur "
                "(1 000 F par jour et par véhicule sous contrat) est déduite des recettes : elle est financée par la part de TamCar.")
    S["A29"] = ("Non inclus : frais FeexPay (1,7 % payés par le client en plus), frais Mobile Money des retraits chauffeur (à décider), forfait internet des chauffeurs (idée non décidée), "
                "charges sociales, entretien (à la charge du conducteur), bonus de performance (neutre aux planchers). Le manque au rachat NSIA est compté (estimation à confirmer).")
    S["A30"] = ("Cellules en bleu sur fond jaune (feuille Hypothèses) : à modifier ; tout le reste se recalcule. Le taux d'atteinte des planchers et les véhicules de tiers "
                "se règlent aux lignes 94 à 102. La feuille Seuils garde le résultat de simulations (valeurs figées, avec leur date).")
    wb.save(dest)
    return dest


if __name__ == "__main__":
    print(construire())
